"""
Tests for the three concrete sectioned-workbook processors: BNMRKProcessor,
AEXPUProcessor and QEXPUProcessor.

The sheet grammar itself is covered by test_sectioned_sheet.py. What is tested
here is everything the concrete subclasses add: glob discovery and its depth,
the filename-token separation between annual and quarterly EXPU files, the
table -> sheet mapping, the path-derived metadata columns (including their NULL
cases), the missing Table 6 delivery, and the incremental / full-refresh
contracts.

Every fixture is a synthetic openpyxl workbook written into tmp_path in a tree
mirroring the real organised file store. No real ACO data is used anywhere.

Fixture sheets deliberately carry content unique to each sheet — a section
label and a row label naming the table they belong to. Sharing one payload
across sheets would make the whole table -> sheet mapping untestable: every
pattern could point at the wrong sheet and the suite would stay green.
"""

from datetime import date
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from mssp_pipeline.processing.exporters.duckdb_exporter import DuckDBExporter
from mssp_pipeline.processing.processors.aexpu_processor import AEXPUProcessor
from mssp_pipeline.processing.processors.bnmrk_processor import BNMRKProcessor
from mssp_pipeline.processing.processors.qexpu_processor import QEXPUProcessor

from .conftest import ACO_ID

# ---------------------------------------------------------------------------
# Synthetic sheets
# ---------------------------------------------------------------------------


def bnmrk_matrix(tag: str, extra_rows=()) -> list:
    """A BNMRK data sheet whose section and row labels name the table it feeds."""
    return [
        [f"Table {tag}", None, None, None],
        ["Shared Savings Program Historical Benchmark Report", None, None, None],
        [f"Table {tag} Calculation", "BY1", "BY2", "Benchmark"],
        [f"[L] Table {tag} Section", None, None, None],
        [f"Row Unique To Table {tag}", "1", "2", f"value-for-table-{tag}"],
        *extra_rows,
    ]


# Table 1 additionally carries the real report's ground-truth benchmark values.
BNMRK_TABLE_1 = bnmrk_matrix(
    "1",
    extra_rows=[
        ["ESRD", "1", "2", "174011.264934806"],
        ["Disabled", "3", "4", "29893.6862970591"],
    ],
)

# Table 6 has its own shape: a decorative single-cell banner above the header
# row, and a bracketed [F] section instead of [L].
BNMRK_TABLE_6 = [
    ["Table 6", None, None],
    ["Accountable Care Prospective Trend", None, None],
    [None, None, "Performance Year"],
    ["Table 6 Calculation", "PY1", "PY2"],
    ["[F] Table 6 Section", None, None],
    ["Row Unique To Table 6", "x", "y"],
    ["ESRD", "1.04962357443411", "1.09"],
]

BNMRK_SHEET_BY_TABLE = {
    "bnmrk_table_1": ("Table 1 - Historical Benchmark", "1", BNMRK_TABLE_1),
    "bnmrk_table_1a": ("Table 1A - Regional Adjustment", "1A", bnmrk_matrix("1A")),
    "bnmrk_table_1b": ("Table 1B - Prior Savings Adj", "1B", bnmrk_matrix("1B")),
    "bnmrk_table_1c": ("Table 1C - Health Equity Adj", "1C", bnmrk_matrix("1C")),
    "bnmrk_table_2": ("Table 2 - Trend Factor", "2", bnmrk_matrix("2")),
    "bnmrk_table_3": ("Table 3 - Truncation", "3", bnmrk_matrix("3")),
    "bnmrk_table_4": ("Table 4 - Renormalization", "4", bnmrk_matrix("4")),
    "bnmrk_table_5": ("Table 5 - Reg Adj Weight", "5", bnmrk_matrix("5")),
    "bnmrk_table_6": ("Table 6 - ACPT", "6", BNMRK_TABLE_6),
}

# The real "Parameters" tab: key/value, no header row, blank column B = section.
PARAMETERS = [
    ["Parameters", None],
    ["Table of Contents", None],
    ["ACO Track", "BASIC Level A"],
    ["Claims-Based Beneficiary Assignment Window", None],
    ["Benchmark Year 1 (BY1)", "01/01/2022 - 12/31/2022"],
    ["Benchmark Year 2 (BY2)", "01/01/2023 - 12/31/2023"],
]


def eu_report(tag: str) -> list:
    return [
        [f"Table {tag}", None, None],
        ["Aggregate Expenditure and Utilization Report", None, None],
        [f"Table {tag} Category", "ACO", "National"],
        [f"[A] Table {tag} Section", None, None],
        [f"Row Unique To Table {tag}", "100", "200"],
    ]


def snf_report(tag: str) -> list:
    return [
        [f"Table {tag}", None, None, None],
        ["Skilled Nursing Facility (SNF) Report", None, None, None],
        [None, "ACO-Specific", "Affiliated SNFs", None],
        [None, None, "Total", "With Prior Stay"],
        [f"Table {tag} Section", None, None, None],
        [f"Row Unique To Table {tag}", "439", "-", "12"],
    ]


def truncation(tag: str) -> list:
    return [
        [f"Table {tag}", None, None],
        ["Expenditure Truncation", None, None],
        [f"Table {tag} Category", "Threshold", "Truncated"],
        [f"[T] Table {tag} Section", None, None],
        [f"Row Unique To Table {tag}", "1", "2"],
    ]


REGIONAL_EXPENDITURES = [
    ["Table 2", None, None],
    ["Regional Expenditures", None, None],
    [None, "Benchmark Year 3", "Q1"],
    ["Regional Expenditures ($)", None, None],
    ["Row Unique To Table 2", "125734", "132431"],
    ["National Weight", None, None],
    ["Row Unique To Table 2", "0.11", "0.12"],
    ["Regional Weight", None, None],
    ["Row Unique To Table 2", "0.21", "0.22"],
]

# The BY1 and BY2 annual workbooks carry COVID-excluding variants of Tables 1
# and 4; BY3 does not. Verified across all six real AEXPU deliveries
# (Y2022/Y2023/Y2024 x two bundles), so AEXPU_TABLE_1A / _4A are legitimately
# empty for the most recent benchmark year.
BENCHMARK_YEARS_WITH_COVID_VARIANTS = ("2022", "2023")


def bnmrk_sheets(with_table_6: bool = True) -> dict:
    sheets = {
        "Cover": [["Cover page"]],
        "TOC": [["Table of Contents"]],
        "Glossary": [["Term", "Definition"]],
        "Parameters": PARAMETERS,
    }
    for table, (sheet_name, _tag, rows) in BNMRK_SHEET_BY_TABLE.items():
        if table == "bnmrk_table_6" and not with_table_6:
            continue
        sheets[sheet_name] = rows
    return sheets


def aexpu_sheets(benchmark_year: str) -> dict:
    sheets = {
        "Cover": [["Cover page"]],
        "TOC": [["Table of Contents"]],
        "Glossary": [["Term", "Definition"]],
        "Parameters": PARAMETERS,
        "Table_1-Aggregate_EU_Report": eu_report("1"),
        "Table_3-SNF_Report": snf_report("3"),
        "Table_4-Exp_Truncation": truncation("4"),
    }
    if benchmark_year in BENCHMARK_YEARS_WITH_COVID_VARIANTS:
        sheets["Table_1A-EU_Excluding_COVID"] = eu_report("1A")
        sheets["Table_4A-Ex_Trc_Excluding_COVID"] = truncation("4A")
    return sheets


def qexpu_sheets() -> dict:
    return {
        "Cover": [["Cover page"]],
        "TOC": [["Table of Contents"]],
        "Glossary": [["Term", "Definition"]],
        "Parameters": PARAMETERS,
        "Table_1-Aggregate_EU_Report": eu_report("1"),
        "Table_2-Regional_Expenditures": REGIONAL_EXPENDITURES,
        "Table_3-SNF_Report": snf_report("3"),
    }


# ---------------------------------------------------------------------------
# File-store builders
# ---------------------------------------------------------------------------


def write_workbook(path: Path, sheets: dict) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(sheet_name)
        for row in rows:
            ws.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def make_bnmrk_bundle(
    raw_dir: Path,
    year: str = "2025",
    code: str = "01",
    date_str: str = "259999",
    time_str: str = "1111111",
    with_table_6: bool = True,
    benchmark_years=("2022", "2023", "2024"),
    aco: str = ACO_ID,
) -> Path:
    """Build one BNMRK bundle:

        raw_dir/T0000/{year}/{code}/P.T0000.ACO.BNMRK.D{date}.T{time}/
            P.T0000.ACO.BNMRK.D{date}.T{time}.xlsx        <- BNMRKProcessor
            P.T0000.ACO.AEXPU.Y2022.D{date}.T{time}.xlsx  <- AEXPUProcessor
            ... one per benchmark year ...
            P.T0000.ACO.AASR.Y2024.D{date}.T{time}.xlsx   <- decoy, matched by neither
    """
    stamp = f"D{date_str}.T{time_str}"
    bundle = raw_dir / aco / year / code / f"P.{aco}.ACO.BNMRK.{stamp}"
    write_workbook(
        bundle / f"P.{aco}.ACO.BNMRK.{stamp}.xlsx", bnmrk_sheets(with_table_6)
    )
    for benchmark_year in benchmark_years:
        write_workbook(
            bundle / f"P.{aco}.ACO.AEXPU.Y{benchmark_year}.{stamp}.xlsx",
            aexpu_sheets(benchmark_year),
        )
    write_workbook(bundle / f"P.{aco}.ACO.AASR.Y2024.{stamp}.xlsx", qexpu_sheets())
    return bundle


def make_qexpu_bundle(
    raw_dir: Path,
    year: str = "2026",
    code: str = "02",
    quarter: str = "2026Q1",
    date_str: str = "269999",
    time_str: str = "0100000",
    aco: str = ACO_ID,
) -> Path:
    stamp = f"D{date_str}.T{time_str}"
    bundle = raw_dir / aco / year / code / f"P.{aco}.ACO.QEXPU.{stamp}"
    write_workbook(bundle / f"P.{aco}.ACO.QEXPU.{quarter}.{stamp}.xlsx", qexpu_sheets())
    return bundle


def config_for(raw_dir, aco: str = ACO_ID, full_refresh: bool = True, **overrides):
    settings = dict(
        ACO_ID=aco,
        FILE_STORE=str(raw_dir),
        FULL_REFRESH=full_refresh,
        PROCESS_BATCH_SIZE_DEFAULT=25,
        PROCESS_BATCH_SIZE_EXPU=25,
    )
    settings.update(overrides)
    return SimpleNamespace(**settings)


# ---------------------------------------------------------------------------
# Runner
# ---------------------------------------------------------------------------


def run(processor_cls, session, config):
    """Run one processor into DuckDB's raw_data schema; return the connection."""
    exporter = DuckDBExporter(schema="raw_data", full_refresh=config.FULL_REFRESH)
    processor_cls(session, exporter, config).run()
    return session.connection


def counts(conn) -> dict:
    return {
        table: conn.execute(f"SELECT COUNT(*) FROM raw_data.{table}").fetchone()[0]
        for (table,) in conn.execute(
            "SELECT table_name FROM information_schema.tables "
            "WHERE table_schema = 'raw_data'"
        ).fetchall()
    }


# ===========================================================================
# Table -> sheet mapping
# ===========================================================================


@pytest.mark.parametrize("table,tag", [(t, v[1]) for t, v in BNMRK_SHEET_BY_TABLE.items()])
def test_each_bnmrk_table_holds_its_own_sheet(
    test_session, test_config, raw_dir, table, tag
):
    """Every BNMRK table must contain the sheet its pattern names — and only it.

    This is the test that guards the whole table -> sheet mapping. Each fixture
    sheet carries a section label and a row label unique to it, so pointing
    BNMRK_TABLE_2 at 'Table 3 - %' fails here rather than passing silently.
    """
    make_bnmrk_bundle(raw_dir)
    conn = run(BNMRKProcessor, test_session, test_config)

    assert conn.execute(
        f"SELECT DISTINCT section_label FROM raw_data.{table}"
    ).fetchall() == [(f"Table {tag} Section",)]
    assert conn.execute(
        f"SELECT COUNT(*) FROM raw_data.{table} "
        f"WHERE row_label = 'Row Unique To Table {tag}'"
    ).fetchone()[0] > 0


def test_each_aexpu_table_holds_its_own_sheet(test_session, test_config, raw_dir):
    """Same guard for AEXPU, whose patterns rely on the '\\_' LIKE escape."""
    make_bnmrk_bundle(raw_dir)
    conn = run(AEXPUProcessor, test_session, test_config)

    for table, tag in (("aexpu_table_1", "1"), ("aexpu_table_1a", "1A"),
                       ("aexpu_table_3", "3"), ("aexpu_table_4", "4"),
                       ("aexpu_table_4a", "4A")):
        assert conn.execute(
            f"SELECT DISTINCT row_label FROM raw_data.{table}"
        ).fetchall() == [(f"Row Unique To Table {tag}",)], table


def test_each_qexpu_table_holds_its_own_sheet(test_session, test_config, raw_dir):
    """Same guard for QEXPU, which is the only family with a Table 2."""
    make_qexpu_bundle(raw_dir)
    conn = run(QEXPUProcessor, test_session, test_config)

    for table, tag in (("qexpu_table_1", "1"), ("qexpu_table_2", "2"),
                       ("qexpu_table_3", "3")):
        assert conn.execute(
            f"SELECT DISTINCT row_label FROM raw_data.{table}"
        ).fetchall() == [(f"Row Unique To Table {tag}",)], table


# ===========================================================================
# BNMRK
# ===========================================================================


def test_bnmrk_glob_finds_only_bnmrk_workbooks(test_session, test_config, raw_dir):
    """The '.BNMRK.' filename token excludes the AEXPU and AASR files beside it."""
    make_bnmrk_bundle(raw_dir)
    conn = run(BNMRKProcessor, test_session, test_config)

    assert conn.execute(
        "SELECT DISTINCT file_name FROM raw_data.bnmrk_table_1"
    ).fetchall() == [(f"P.{ACO_ID}.ACO.BNMRK.D259999.T1111111.xlsx",)]


def test_bnmrk_glob_depth_is_year_code_bundle_file(test_session, test_config, raw_dir):
    """Discovery is pinned to {ACO}/{YEAR}/{CODE}/{BUNDLE}/{file} — no other depth.

    Matches the depth EXPUProcessor globbed against the real S3 store; a
    workbook one level shallower or one level deeper is a layout the organised
    store does not produce and must not be swept up.
    """
    make_bnmrk_bundle(raw_dir)
    stamp = "D259999.T2222222"
    # One level too shallow: {ACO}/{YEAR}/{BUNDLE}/{file}.
    write_workbook(
        raw_dir / ACO_ID / "2025" / f"P.{ACO_ID}.ACO.BNMRK.{stamp}"
        / f"P.{ACO_ID}.ACO.BNMRK.{stamp}.xlsx",
        bnmrk_sheets(),
    )
    # One level too deep.
    write_workbook(
        raw_dir / ACO_ID / "2025" / "01" / "extra" / f"P.{ACO_ID}.ACO.BNMRK.{stamp}"
        / f"P.{ACO_ID}.ACO.BNMRK.{stamp}.xlsx",
        bnmrk_sheets(),
    )

    conn = run(BNMRKProcessor, test_session, test_config)
    assert conn.execute(
        "SELECT DISTINCT submission_id FROM raw_data.bnmrk_table_1"
    ).fetchall() == [("D259999.T1111111",)]


def test_bnmrk_metadata_columns(test_session, test_config, raw_dir):
    """PERFORMANCE_YEAR from the path; SUBMISSION_ID and PERIOD from the filename.

    BENCHMARK_YEAR is not a BNMRK concept and stays NULL, and FILE_DATE is NULL
    because 'D259999' is a placeholder rather than a real YYMMDD stamp.
    """
    make_bnmrk_bundle(raw_dir)
    conn = run(BNMRKProcessor, test_session, test_config)

    assert conn.execute(
        "SELECT DISTINCT aco_id, performance_year, benchmark_year, submission_id, "
        "period, file_date FROM raw_data.bnmrk_table_1"
    ).fetchall() == [
        (ACO_ID, "2025", None, "D259999.T1111111", "D259999.T1111111", None)
    ]


def test_bnmrk_file_date_parses_a_real_submission_stamp(
    test_session, test_config, raw_dir
):
    """A genuine D<YYMMDD> stamp becomes a date; only placeholders go NULL."""
    make_bnmrk_bundle(raw_dir, date_str="250709", time_str="1405028")
    conn = run(BNMRKProcessor, test_session, test_config)

    assert conn.execute(
        "SELECT DISTINCT file_date, submission_id FROM raw_data.bnmrk_table_1"
    ).fetchall() == [(date(2025, 7, 9), "D250709.T1405028")]


@pytest.mark.parametrize("date_str", ["259999", "250230"])
def test_bnmrk_file_date_is_null_for_an_unparseable_stamp(
    test_session, test_config, raw_dir, date_str
):
    """Two distinct failures: an impossible month, and an impossible day.

    'D259999' is CMS's placeholder (month 99); 'D250230' is 30 February — a real
    month with a day that does not exist in it. Both must yield NULL rather than
    a rolled-over or clamped date, and neither may cause the file to be skipped:
    the submission stamp is still well-formed, so the delivery is still
    identifiable.
    """
    make_bnmrk_bundle(raw_dir, date_str=date_str)
    conn = run(BNMRKProcessor, test_session, test_config)

    assert conn.execute(
        "SELECT DISTINCT file_date, submission_id FROM raw_data.bnmrk_table_1"
    ).fetchall() == [(None, f"D{date_str}.T1111111")]


def test_bnmrk_performance_year_is_null_when_the_segment_is_not_a_year(
    test_session, test_config, raw_dir
):
    """A path that does not carry a 4-digit year yields NULL, not a bad guess."""
    make_bnmrk_bundle(raw_dir, year="PY25")
    conn = run(BNMRKProcessor, test_session, test_config)

    assert conn.execute(
        "SELECT DISTINCT performance_year FROM raw_data.bnmrk_table_1"
    ).fetchall() == [(None,)]


def test_performance_year_when_file_store_itself_ends_in_the_aco_id(
    test_session, raw_dir
):
    """'.../T0000/T0000/2025/...' — the nearer occurrence is the right one."""
    store = raw_dir / ACO_ID
    make_bnmrk_bundle(store)
    conn = run(BNMRKProcessor, test_session, config_for(store))

    assert conn.execute(
        "SELECT DISTINCT performance_year FROM raw_data.bnmrk_table_1"
    ).fetchall() == [("2025",)]


def test_performance_year_when_the_code_segment_is_the_aco_id(test_session, raw_dir):
    """'.../T0000/2025/T0000/...' — the nearer occurrence is followed by a bundle
    directory, so the scan must keep going rather than give up and return NULL."""
    make_bnmrk_bundle(raw_dir, code=ACO_ID)
    conn = run(BNMRKProcessor, test_session, config_for(raw_dir))

    assert conn.execute(
        "SELECT DISTINCT performance_year FROM raw_data.bnmrk_table_1"
    ).fetchall() == [("2025",)]


def test_trailing_slash_file_store_yields_identical_file_paths(
    test_session, raw_dir
):
    """A FILE_STORE written with a trailing slash must not shift FILE_PATH.

    FILE_PATH is the dedup key, so any difference — a doubled separator, say —
    would make every file look new on the next run.
    """
    make_bnmrk_bundle(raw_dir)

    plain = run(BNMRKProcessor, test_session, config_for(raw_dir))
    without = plain.execute(
        "SELECT DISTINCT file_path FROM raw_data.bnmrk_table_1"
    ).fetchall()

    trailing = run(BNMRKProcessor, test_session, config_for(f"{raw_dir}/"))
    with_slash = trailing.execute(
        "SELECT DISTINCT file_path FROM raw_data.bnmrk_table_1"
    ).fetchall()

    assert with_slash == without
    assert "//" not in with_slash[0][0]


def test_a_file_without_a_submission_stamp_is_skipped(
    test_session, test_config, raw_dir, capsys
):
    """A stray '...T1111111.backup.xlsx' is loudly ignored, not silently loaded.

    The glob cannot exclude it — 'T*' matches '1111111.backup' — and loading it
    would give it NULL SUBMISSION_ID / PERIOD / FILE_DATE. For BNMRK the
    submission id is the only thing separating deliveries, so two such files
    would merge into one indistinguishable identity.
    """
    bundle = make_bnmrk_bundle(raw_dir)
    write_workbook(
        bundle / f"P.{ACO_ID}.ACO.BNMRK.D259999.T1111111.backup.xlsx", bnmrk_sheets()
    )

    conn = run(BNMRKProcessor, test_session, test_config)

    assert conn.execute(
        "SELECT DISTINCT file_name FROM raw_data.bnmrk_table_1"
    ).fetchall() == [(f"P.{ACO_ID}.ACO.BNMRK.D259999.T1111111.xlsx",)]
    assert conn.execute(
        "SELECT COUNT(*) FROM raw_data.bnmrk_table_1 WHERE submission_id IS NULL"
    ).fetchone()[0] == 0
    assert "backup.xlsx" in capsys.readouterr().out


def test_two_bnmrk_deliveries_are_distinguished_by_submission_id(
    test_session, test_config, raw_dir
):
    """March / June / October deliveries differ only in the D/T stamp."""
    make_bnmrk_bundle(raw_dir, time_str="1111111")
    make_bnmrk_bundle(raw_dir, time_str="0000000", with_table_6=False)
    conn = run(BNMRKProcessor, test_session, test_config)

    assert sorted(
        r[0] for r in conn.execute(
            "SELECT DISTINCT submission_id FROM raw_data.bnmrk_table_1"
        ).fetchall()
    ) == ["D259999.T0000000", "D259999.T1111111"]
    assert conn.execute(
        "SELECT COUNT(*) FROM raw_data.bnmrk_table_1 WHERE file_date IS NOT NULL"
    ).fetchone()[0] == 0


@pytest.mark.parametrize("batch_size", [1, 25])
def test_bnmrk_delivery_without_table_6_does_not_fail_the_run(
    test_session, raw_dir, batch_size
):
    """'Table 6 - ACPT' is absent from the March preliminary — at any batch size.

    At batch size 1 the Table-6-less delivery fills a whole batch, which is the
    case that used to fail the table outright. PROCESS_BATCH_SIZE_DEFAULT is set
    to a different value on purpose, so a processor that has fallen out of
    base.py's batch-size map would silently run at the default instead.
    """
    config = config_for(
        raw_dir,
        PROCESS_BATCH_SIZE_DEFAULT=99,
        PROCESS_BATCH_SIZE_EXPU=batch_size,
    )
    make_bnmrk_bundle(raw_dir, time_str="0000000", with_table_6=False)
    make_bnmrk_bundle(raw_dir, time_str="1111111", with_table_6=True)

    processor = BNMRKProcessor(test_session, DuckDBExporter(), config)
    assert processor._batch_size_for() == batch_size

    conn = run(BNMRKProcessor, test_session, config)

    # Every other table loads from both deliveries...
    assert conn.execute(
        "SELECT COUNT(DISTINCT submission_id) FROM raw_data.bnmrk_table_1"
    ).fetchone()[0] == 2
    # ...while Table 6 loads only from the delivery that has it.
    assert conn.execute(
        "SELECT DISTINCT submission_id FROM raw_data.bnmrk_table_6"
    ).fetchall() == [("D259999.T1111111",)]
    assert conn.execute(
        "SELECT value_text FROM raw_data.bnmrk_table_6 "
        "WHERE section_code = 'F' AND row_label = 'ESRD' AND column_label = 'PY1'"
    ).fetchone()[0] == "1.04962357443411"


@pytest.mark.parametrize(
    "processor_cls", [BNMRKProcessor, AEXPUProcessor, QEXPUProcessor]
)
def test_all_three_processors_use_the_expu_batch_size_knob(
    test_session, raw_dir, processor_cls
):
    """base.py keys its batch-size map on the class name — all three must be in it.

    MSSP_PROCESS_BATCH_SIZE_EXPU is the pre-existing env knob these three
    inherited from the retired EXPUProcessor; dropping a key would silently fall
    back to the default.
    """
    config = config_for(
        raw_dir, PROCESS_BATCH_SIZE_DEFAULT=99, PROCESS_BATCH_SIZE_EXPU=7
    )
    processor = processor_cls(test_session, DuckDBExporter(), config)
    assert processor._batch_size_for() == 7


def test_bnmrk_parameters_is_read_as_key_values(test_session, test_config, raw_dir):
    """The Parameters tab has no header row; synthetic labels carry it."""
    make_bnmrk_bundle(raw_dir)
    conn = run(BNMRKProcessor, test_session, test_config)

    assert conn.execute(
        "SELECT value_text FROM raw_data.bnmrk_parameters WHERE row_label = 'ACO Track'"
    ).fetchone()[0] == "BASIC Level A"
    assert conn.execute(
        "SELECT DISTINCT section_label FROM raw_data.bnmrk_parameters "
        "WHERE row_label LIKE 'Benchmark Year%'"
    ).fetchall() == [("Claims-Based Beneficiary Assignment Window",)]


# ===========================================================================
# AEXPU
# ===========================================================================


def test_aexpu_glob_finds_files_inside_the_bnmrk_bundle(
    test_session, test_config, raw_dir
):
    """Annual EXPU workbooks live in a BNMRK-named bundle directory.

    That is exactly why the glob keys on the '.AEXPU.' filename token: a
    directory-name match would find none of them.
    """
    make_bnmrk_bundle(raw_dir)
    conn = run(AEXPUProcessor, test_session, test_config)

    assert sorted(
        r[0] for r in conn.execute(
            "SELECT DISTINCT file_name FROM raw_data.aexpu_table_1"
        ).fetchall()
    ) == [
        f"P.{ACO_ID}.ACO.AEXPU.Y2022.D259999.T1111111.xlsx",
        f"P.{ACO_ID}.ACO.AEXPU.Y2023.D259999.T1111111.xlsx",
        f"P.{ACO_ID}.ACO.AEXPU.Y2024.D259999.T1111111.xlsx",
    ]


def test_aexpu_glob_ignores_quarterly_files(test_session, test_config, raw_dir):
    """'.AEXPU.' must not match '.QEXPU.' — the old '*EXPU*' glob matched both."""
    make_bnmrk_bundle(raw_dir)
    make_qexpu_bundle(raw_dir)
    conn = run(AEXPUProcessor, test_session, test_config)

    assert conn.execute(
        "SELECT COUNT(*) FROM raw_data.aexpu_table_1 WHERE file_name LIKE '%QEXPU%'"
    ).fetchone()[0] == 0
    assert conn.execute(
        "SELECT COUNT(DISTINCT file_name) FROM raw_data.aexpu_table_1"
    ).fetchone()[0] == 3


def test_aexpu_metadata_columns(test_session, test_config, raw_dir):
    """BENCHMARK_YEAR / PERIOD / FILE_DATE all come from the Y<year> token."""
    make_bnmrk_bundle(raw_dir)
    conn = run(AEXPUProcessor, test_session, test_config)

    assert conn.execute(
        "SELECT DISTINCT benchmark_year, period, file_date, performance_year, "
        "submission_id, aco_id FROM raw_data.aexpu_table_1 "
        "WHERE benchmark_year = '2024'"
    ).fetchall() == [
        ("2024", "Y2024", date(2024, 12, 31), "2025", "D259999.T1111111", ACO_ID)
    ]


def test_three_aexpu_files_in_one_bundle_give_three_benchmark_years(
    test_session, test_config, raw_dir
):
    """BY1/BY2/BY3 land in one table, told apart by BENCHMARK_YEAR alone.

    Their PERFORMANCE_YEAR and SUBMISSION_ID are identical — the whole point of
    the column.
    """
    make_bnmrk_bundle(raw_dir)
    conn = run(AEXPUProcessor, test_session, test_config)

    assert conn.execute(
        "SELECT DISTINCT benchmark_year, period, file_date "
        "FROM raw_data.aexpu_table_1 ORDER BY 1"
    ).fetchall() == [
        ("2022", "Y2022", date(2022, 12, 31)),
        ("2023", "Y2023", date(2023, 12, 31)),
        ("2024", "Y2024", date(2024, 12, 31)),
    ]
    assert conn.execute(
        "SELECT COUNT(DISTINCT performance_year || submission_id) "
        "FROM raw_data.aexpu_table_1"
    ).fetchone()[0] == 1


def test_aexpu_benchmark_year_is_anchored_to_the_family_token(test_session, raw_dir):
    """A 'Y####'-shaped ACO id must not be mistaken for the benchmark year.

    'P.Y2024.ACO.AEXPU.Y2022.D...' has two '.Y####.' tokens; an unanchored
    search takes the leftmost and reports benchmark year 2024 for the BY1
    workbook. BENCHMARK_YEAR is the only thing separating the three workbooks of
    a bundle, so that would collapse them.
    """
    aco = "Y2024"
    make_bnmrk_bundle(raw_dir, aco=aco)
    conn = run(AEXPUProcessor, test_session, config_for(raw_dir, aco=aco))

    assert conn.execute(
        "SELECT DISTINCT benchmark_year, period, file_date "
        "FROM raw_data.aexpu_table_1 ORDER BY 1"
    ).fetchall() == [
        ("2022", "Y2022", date(2022, 12, 31)),
        ("2023", "Y2023", date(2023, 12, 31)),
        ("2024", "Y2024", date(2024, 12, 31)),
    ]


def test_aexpu_without_a_year_token_gets_null_benchmark_metadata(
    test_session, test_config, raw_dir
):
    """No Y<year> in the filename means no benchmark year to report.

    The file is still identifiable — its submission stamp parses — so it loads,
    with the three year-derived columns NULL rather than guessed.
    """
    bundle = make_bnmrk_bundle(raw_dir, benchmark_years=())
    write_workbook(
        bundle / f"P.{ACO_ID}.ACO.AEXPU.D259999.T1111111.xlsx", aexpu_sheets("2024")
    )
    conn = run(AEXPUProcessor, test_session, test_config)

    assert conn.execute(
        "SELECT DISTINCT benchmark_year, period, file_date, submission_id "
        "FROM raw_data.aexpu_table_1"
    ).fetchall() == [(None, None, None, "D259999.T1111111")]


def test_aexpu_covid_variants_do_not_bleed_into_the_base_tables(
    test_session, test_config, raw_dir
):
    r"""'Table\_1-%' must not swallow 'Table_1A-EU_Excluding_COVID', or vice versa.

    Both sheets now have their own destination, so a pattern that matched the
    wrong one would double one table and empty the other rather than merely
    adding noise.
    """
    make_bnmrk_bundle(raw_dir)
    conn = run(AEXPUProcessor, test_session, test_config)

    for table, tag in (("aexpu_table_1", "1"), ("aexpu_table_1a", "1A"),
                       ("aexpu_table_4", "4"), ("aexpu_table_4a", "4A")):
        assert conn.execute(
            f"SELECT DISTINCT row_label FROM raw_data.{table}"
        ).fetchall() == [(f"Row Unique To Table {tag}",)], table

    # Every AEXPU file contributes the same number of Table 1 rows.
    assert len(set(conn.execute(
        "SELECT COUNT(*) FROM raw_data.aexpu_table_1 GROUP BY file_name"
    ).fetchall())) == 1


def test_aexpu_covid_variants_load_only_from_the_workbooks_that_have_them(
    test_session, test_config, raw_dir
):
    """BY1 and BY2 carry Tables 1A/4A; BY3 does not — and that is not an error.

    The base tables load from all three workbooks, the variants from two. A
    downstream model must expect AEXPU_TABLE_1A / _4A to be legitimately empty
    for the most recent benchmark year rather than treat it as missing data.
    """
    make_bnmrk_bundle(raw_dir)
    conn = run(AEXPUProcessor, test_session, test_config)

    assert sorted(r[0] for r in conn.execute(
        "SELECT DISTINCT benchmark_year FROM raw_data.aexpu_table_1"
    ).fetchall()) == ["2022", "2023", "2024"]

    for table in ("aexpu_table_1a", "aexpu_table_4a"):
        assert sorted(r[0] for r in conn.execute(
            f"SELECT DISTINCT benchmark_year FROM raw_data.{table}"
        ).fetchall()) == ["2022", "2023"], table


@pytest.mark.parametrize("table", ["aexpu_table_1a", "aexpu_table_4a"])
def test_aexpu_covid_variant_metadata(test_session, test_config, raw_dir, table):
    """The variants carry the same path-derived metadata as every other table."""
    make_bnmrk_bundle(raw_dir)
    conn = run(AEXPUProcessor, test_session, test_config)

    assert conn.execute(
        f"SELECT DISTINCT benchmark_year, period, file_date, performance_year, "
        f"submission_id, aco_id FROM raw_data.{table} ORDER BY 1"
    ).fetchall() == [
        ("2022", "Y2022", date(2022, 12, 31), "2025", "D259999.T1111111", ACO_ID),
        ("2023", "Y2023", date(2023, 12, 31), "2025", "D259999.T1111111", ACO_ID),
    ]


def test_aexpu_bundle_of_only_variant_less_workbooks_is_not_a_failure(
    test_session, raw_dir
):
    """A whole batch of BY3-only workbooks contributes 0 rows without failing.

    At batch size 1 the BY3 workbook fills its own batch, which is the case the
    per-batch missing-sheet path exists for.
    """
    config = config_for(raw_dir, PROCESS_BATCH_SIZE_DEFAULT=99, PROCESS_BATCH_SIZE_EXPU=1)
    make_bnmrk_bundle(raw_dir)
    conn = run(AEXPUProcessor, test_session, config)

    assert conn.execute("SELECT COUNT(*) FROM raw_data.aexpu_table_1a").fetchone()[0] > 0
    assert conn.execute(
        "SELECT COUNT(*) FROM raw_data.aexpu_table_1a WHERE benchmark_year = '2024'"
    ).fetchone()[0] == 0


def test_aexpu_has_no_table_2(test_session, test_config, raw_dir):
    """Regional expenditures are quarterly-only; there is no AEXPU_TABLE_2 def."""
    make_bnmrk_bundle(raw_dir)
    conn = run(AEXPUProcessor, test_session, test_config)

    tables = set(counts(conn))
    assert "aexpu_table_2" not in tables
    assert {"aexpu_table_1", "aexpu_table_1a", "aexpu_table_3", "aexpu_table_4",
            "aexpu_table_4a", "aexpu_parameters"} <= tables


# ===========================================================================
# QEXPU
# ===========================================================================


def test_qexpu_glob_ignores_annual_files(test_session, test_config, raw_dir):
    """The reverse separation: '.QEXPU.' must not match '.AEXPU.'."""
    make_bnmrk_bundle(raw_dir)
    make_qexpu_bundle(raw_dir)
    conn = run(QEXPUProcessor, test_session, test_config)

    assert conn.execute(
        "SELECT DISTINCT file_name FROM raw_data.qexpu_table_1"
    ).fetchall() == [(f"P.{ACO_ID}.ACO.QEXPU.2026Q1.D269999.T0100000.xlsx",)]


def test_qexpu_metadata_columns(test_session, test_config, raw_dir):
    """PERIOD is the literal quarter, FILE_DATE its last calendar day."""
    make_qexpu_bundle(raw_dir)
    conn = run(QEXPUProcessor, test_session, test_config)

    assert conn.execute(
        "SELECT DISTINCT aco_id, performance_year, benchmark_year, submission_id, "
        "period, file_date FROM raw_data.qexpu_table_1"
    ).fetchall() == [
        (ACO_ID, "2026", None, "D269999.T0100000", "2026Q1", date(2026, 3, 31))
    ]


@pytest.mark.parametrize(
    "quarter,expected",
    [("2026Q1", date(2026, 3, 31)), ("2026Q2", date(2026, 6, 30)),
     ("2026Q3", date(2026, 9, 30)), ("2026Q4", date(2026, 12, 31))],
)
def test_qexpu_file_date_is_the_quarter_end(
    test_session, test_config, raw_dir, quarter, expected
):
    """Every quarter maps to its own last day, December included."""
    make_qexpu_bundle(raw_dir, quarter=quarter)
    conn = run(QEXPUProcessor, test_session, test_config)

    assert conn.execute(
        "SELECT DISTINCT period, file_date FROM raw_data.qexpu_table_1"
    ).fetchall() == [(quarter, expected)]


def test_qexpu_table_2_distinguishes_its_sections(test_session, test_config, raw_dir):
    """Table 2 stacks three blocks under one header; SECTION_LABEL separates them."""
    make_qexpu_bundle(raw_dir)
    conn = run(QEXPUProcessor, test_session, test_config)

    assert sorted(
        r[0] for r in conn.execute(
            "SELECT DISTINCT section_label FROM raw_data.qexpu_table_2"
        ).fetchall()
    ) == ["National Weight", "Regional Expenditures ($)", "Regional Weight"]
    assert conn.execute(
        "SELECT value_text FROM raw_data.qexpu_table_2 "
        "WHERE section_label = 'Regional Weight' AND column_label = 'Q1'"
    ).fetchone()[0] == "0.22"


# ===========================================================================
# Incremental / full refresh
# ===========================================================================


@pytest.mark.parametrize(
    "processor_cls,builder",
    [(BNMRKProcessor, make_bnmrk_bundle),
     (AEXPUProcessor, make_bnmrk_bundle),
     (QEXPUProcessor, make_qexpu_bundle)],
)
def test_second_incremental_run_adds_zero_rows(
    test_session, incremental_config, raw_dir, processor_cls, builder
):
    """The dedup contract: FILE_PATH round-trips, so nothing is re-appended.

    This is the test that catches a _list_source_file_paths() whose FILE_PATH
    normalisation has drifted from _file_metadata_sql() — the failure mode there
    is silent duplication on every single run.
    """
    builder(raw_dir)
    conn = run(processor_cls, test_session, incremental_config)
    before = counts(conn)
    assert any(count > 0 for count in before.values())

    run(processor_cls, test_session, incremental_config)
    assert counts(conn) == before


def test_incremental_run_picks_up_only_the_new_delivery(
    test_session, incremental_config, raw_dir
):
    """A second delivery is appended; the first is not read again."""
    make_bnmrk_bundle(raw_dir, time_str="0000000", with_table_6=False)
    conn = run(BNMRKProcessor, test_session, incremental_config)
    first_rows = conn.execute("SELECT COUNT(*) FROM raw_data.bnmrk_table_1").fetchone()[0]

    make_bnmrk_bundle(raw_dir, time_str="1111111")
    run(BNMRKProcessor, test_session, incremental_config)

    assert conn.execute(
        "SELECT COUNT(*) FROM raw_data.bnmrk_table_1"
    ).fetchone()[0] == first_rows * 2
    assert conn.execute(
        "SELECT submission_id, COUNT(*) FROM raw_data.bnmrk_table_1 "
        "GROUP BY 1 ORDER BY 1"
    ).fetchall() == [
        ("D259999.T0000000", first_rows), ("D259999.T1111111", first_rows)
    ]


def test_full_refresh_replaces_rather_than_appends(
    test_session, test_config, incremental_config, raw_dir
):
    """FULL_REFRESH=True rewrites the table; row counts do not accumulate."""
    make_bnmrk_bundle(raw_dir)
    conn = run(BNMRKProcessor, test_session, incremental_config)
    baseline = counts(conn)

    run(BNMRKProcessor, test_session, test_config)  # FULL_REFRESH=True
    assert counts(conn) == baseline


def test_file_path_matches_the_value_used_for_dedup(
    test_session, test_config, raw_dir
):
    """_list_source_file_paths() and _file_metadata_sql() must agree exactly."""
    make_bnmrk_bundle(raw_dir)
    exporter = DuckDBExporter(schema="raw_data", full_refresh=True)
    processor = BNMRKProcessor(test_session, exporter, test_config)
    listed = {file_path for file_path, _ in processor._list_source_file_paths(None)}
    processor.run()

    emitted = {
        r[0] for r in test_session.connection.execute(
            "SELECT DISTINCT file_path FROM raw_data.bnmrk_table_1"
        ).fetchall()
    }
    assert emitted == listed


def test_missing_file_store_is_not_an_error(test_session, test_config, raw_dir):
    """An empty store means "nothing to do", not a failed run."""
    for processor_cls in (BNMRKProcessor, AEXPUProcessor, QEXPUProcessor):
        run(processor_cls, test_session, test_config)

    assert counts(test_session.connection) == {}
